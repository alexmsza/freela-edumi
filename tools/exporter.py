"""Compilador determinístico dos entregáveis (.csv, .xlsx da Gabi e data_payload.js) com trava de auditoria."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path
from typing import Any

import openpyxl

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from tools.audit import OFFICIAL_SCHEMAS, run_audit
from tools.department_filter import is_clean_departmental_contact


def classify_institution_type(nome_faculdade: str, email: str) -> str:
    """Classifica o tipo da IES exclusivamente para os filtros visuais do dashboard.html."""
    n = (nome_faculdade or "").lower()
    dom = email.split("@")[-1].lower() if "@" in email else ""

    if any(k in n for k in ("escs", "fepecs")) or dom in {"escs.edu.br", "fepecs.edu.br"}:
        return "Pública Distrital (DF)"
    if (
        "instituto federal" in n
        or "cefet" in n
        or re.search(r"\b(if[a-z]{1,12}|cefet)\b", n)
        or re.search(r"(^|\.)(if[a-z-]+|cefet[a-z-]*)\.(edu\.)?br$", dom)
    ):
        return "Instituto Federal (IF)"
    if (
        "universidade federal" in n
        or "universidade de brasília" in n
        or "universidade tecnológica federal" in n
        or any(k in n for k in ("enap", "fiocruz", "ita ", "instituto tecnológico de aeronáutica"))
        or dom.endswith("unb.br")
        or dom in {"enap.gov.br", "fiocruz.br", "ita.br", "utfpr.edu.br", "unifesp.br", "unirio.br", "ufabc.edu.br"}
        or re.search(r"(^|\.)(uf[a-z0-9-]+|unifal-mg|unifei|unila|unilab|unipampa|univasf|unifap|unir|ufrr)\.(edu\.)?br$", dom)
    ):
        return "Universidade Federal"
    if (
        "estadual" in n
        or "fatec" in n
        or re.search(r"\b(usp|unicamp|unesp|uerj|uenf|ueg|uemg|unimontes|upe|uece|uema|uesb|uesc|uefs|uneb|udesc|uel|uem|uepg|unioeste|uea|uerr|unitins|uepb|uern|uncisal|uespi|uepa|uems|unemat)\b", n)
        or re.search(r"(^|\.)(usp|unicamp|unesp|uerj|uenf|ueg|uemg|unimontes|upe|uece|uema|uesb|uesc|uefs|uneb|udesc|uel|uem|uepg|unioeste|uea|uerr|unitins|uepb|uern|uncisal|uespi|uepa|uems|unemat|fatec)\.(edu\.|sp\.gov\.)?br$", dom)
    ):
        return "Universidade Estadual"
    return "Privada / Particular"


def build_deliverables(project_root: Path, strict: bool = True) -> dict[str, Any]:
    """Consolida data_staged/ + data_verified/, executa auditoria e gera CSVs, XLSX e data_payload.js."""
    staged_dir = project_root / "data_staged"
    verified_dir = project_root / "data_verified"
    exports_dir = project_root / "csv_exports"
    exports_dir.mkdir(parents=True, exist_ok=True)

    if not (staged_dir / "01_universidades.csv").exists():
        raise FileNotFoundError("Execute 'uv run edumi ingest-legacy' antes de compilar os entregáveis.")

    def read_staged(fname: str) -> list[dict[str, str]]:
        p = staged_dir / fname
        if not p.exists():
            return []
        with p.open(encoding="utf-8-sig", newline="") as fp:
            return list(csv.DictReader(fp))

    uni_by_email = {r["email"].lower(): r for r in read_staged("01_universidades.csv")}
    ies_file = verified_dir / "ies_scraped.json"
    if ies_file.exists():
        for item in json.loads(ies_file.read_text(encoding="utf-8")):
            em = item["email"].lower().strip()
            uni_by_email[em] = {
                "nome_faculdade": item["nome_faculdade"],
                "nome_departamento": item["nome_departamento"],
                "email": em,
                "uf": item["uf"],
            }

    uni_all = [r for r in uni_by_email.values() if is_clean_departmental_contact(r)]
    uni_all = sorted(uni_all, key=lambda x: (x["uf"] != "DF", x["uf"], x["nome_faculdade"], x["nome_departamento"]))
    uni_df = [r for r in uni_all if r["uf"] == "DF"]

    artistas: list[dict[str, Any]] = []
    influenciadores: list[dict[str, Any]] = []
    social_file = verified_dir / "social_verified.json"
    if social_file.exists():
        for item in json.loads(social_file.read_text(encoding="utf-8")):
            clean_row = {
                "nome": item["nome"],
                "arroba": item["arroba"],
                "quantidade_seguidores": str(item["quantidade_seguidores"]),
                "tipo": item["tipo"],
                "uf": item["uf"],
            }
            if item["categoria"] == "artistas":
                artistas.append(clean_row)
            elif item["categoria"] == "influenciadores":
                influenciadores.append(clean_row)

    artistas.sort(key=lambda x: -int(x["quantidade_seguidores"]))
    influenciadores.sort(key=lambda x: (x["uf"], -int(x["quantidade_seguidores"])))

    midias = read_staged("04_midias.csv")
    parceiros = read_staged("05_parceiros.csv")

    # Grava de volta em staging primeiro para validar antes de publicar em csv_exports
    def write_csv(folder: Path, fname: str, cols: list[str], rows: list[dict[str, Any]]) -> None:
        with (folder / fname).open("w", encoding="utf-8-sig", newline="") as fp:
            w = csv.DictWriter(fp, fieldnames=cols)
            w.writeheader()
            w.writerows(rows)

    write_csv(staged_dir, "01_universidades_DF_marco1.csv", OFFICIAL_SCHEMAS["universidades"], uni_df)
    write_csv(staged_dir, "01_universidades.csv", OFFICIAL_SCHEMAS["universidades"], uni_all)
    write_csv(staged_dir, "02_artistas.csv", OFFICIAL_SCHEMAS["artistas"], artistas)
    write_csv(staged_dir, "03_influenciadores.csv", OFFICIAL_SCHEMAS["influenciadores"], influenciadores)

    audit_res = run_audit(project_root, target_dir=staged_dir)
    if strict and audit_res["total_blockers"] > 0:
        raise RuntimeError(
            f"Build bloqueado pela trava de qualidade: {audit_res['total_blockers']} bloqueador(es) detectado(s). "
            "Consulte reports/AUDITORIA_QUALIDADE.md."
        )

    # Publica CSVs oficiais
    write_csv(exports_dir, "01_universidades_DF_marco1.csv", OFFICIAL_SCHEMAS["universidades"], uni_df)
    write_csv(exports_dir, "01_universidades.csv", OFFICIAL_SCHEMAS["universidades"], uni_all)
    write_csv(exports_dir, "02_artistas.csv", OFFICIAL_SCHEMAS["artistas"], artistas)
    write_csv(exports_dir, "03_influenciadores.csv", OFFICIAL_SCHEMAS["influenciadores"], influenciadores)
    write_csv(exports_dir, "04_midias.csv", OFFICIAL_SCHEMAS["midias"], midias)
    write_csv(exports_dir, "05_parceiros.csv", OFFICIAL_SCHEMAS["parceiros"], parceiros)

    tabs_payload: dict[str, list[dict[str, Any]]] = {
        "universidades": uni_all,
        "artistas": artistas,
        "influenciadores": influenciadores,
        "midias": midias,
        "parceiros": parceiros,
    }

    # Gera a planilha principal '0. Planilha centralizadora.xlsx' (e mantém as cópias auxiliares sincronizadas)
    header_font = Font(bold=True, color="0F172A")
    header_fill = PatternFill(fill_type="solid", fgColor="F1F5F9")
    target_workbooks = (
        "0. Planilha centralizadora.xlsx",
        "0. Planilha centralizadora - Gabi.xlsx",
        "Planilha_Centralizadora_Edumi.xlsx",
    )
    for xlsx_name in target_workbooks:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for tab_name in ("universidades", "artistas", "influenciadores", "midias", "parceiros"):
            ws = wb.create_sheet(title=tab_name)
            cols = OFFICIAL_SCHEMAS[tab_name]
            ws.append(cols)
            ws.freeze_panes = "A2"
            for col_idx in range(1, len(cols) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
            for row in tabs_payload[tab_name]:
                vals = []
                for c in cols:
                    v = row.get(c, "")
                    if c == "quantidade_seguidores" and str(v).isdigit():
                        vals.append(int(v))
                    else:
                        vals.append(v)
                ws.append(vals)
            for col_idx, col_name in enumerate(cols, start=1):
                sample_vals = [len(col_name)] + [len(str(r.get(col_name, ""))) for r in tabs_payload[tab_name][:200]]
                width = min(max(sample_vals) + 4, 58)
                ws.column_dimensions[get_column_letter(col_idx)].width = max(width, 12)
        wb.save(project_root / xlsx_name)

    # Gera data_payload.js com hash de versão e metadado visual tipo_instituicao para dashboard.html
    uni_web = [
        {**r, "tipo_instituicao": classify_institution_type(r["nome_faculdade"], r["email"])}
        for r in uni_all
    ]
    web_payload: dict[str, list[dict[str, Any]]] = {
        "universidades": uni_web,
        "artistas": artistas,
        "influenciadores": influenciadores,
        "midias": midias,
        "parceiros": parceiros,
    }
    raw_json = json.dumps(web_payload, ensure_ascii=False, indent=2)
    payload_hash = hashlib.sha256(raw_json.encode("utf-8")).hexdigest()[:12]
    js_content = (
        f"// Gerado automaticamente por tools/exporter.py — Hash: {payload_hash}\n"
        f"const MASTER_PAYLOAD_VERSION = {json.dumps(payload_hash)};\n"
        f"const MASTER_PAYLOAD = {raw_json};\n"
    )
    (project_root / "data_payload.js").write_text(js_content, encoding="utf-8")

    return {
        "payload_version": payload_hash,
        "counts": {k: len(v) for k, v in tabs_payload.items()},
        "universidades_df": len(uni_df),
    }
