"""Motor de auditoria forense, validação de DNS/MX, filtro anti-bounce e trava anti-alucinação."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import dns.resolver
import openpyxl

from tools.uf_catalog import VALID_UFS

OFFICIAL_SCHEMAS: dict[str, list[str]] = {
    "universidades": ["nome_faculdade", "nome_departamento", "email", "uf"],
    "artistas": ["nome", "arroba", "quantidade_seguidores", "tipo", "uf"],
    "influenciadores": ["nome", "arroba", "quantidade_seguidores", "tipo", "uf"],
    "midias": ["nome", "link", "email", "telefone", "uf", "obs"],
    "parceiros": ["nome", "link", "email", "telefone", "uf", "obs"],
}

CSV_FILE_MAP: dict[str, str] = {
    "universidades_df": "01_universidades_DF_marco1.csv",
    "universidades": "01_universidades.csv",
    "artistas": "02_artistas.csv",
    "influenciadores": "03_influenciadores.csv",
    "midias": "04_midias.csv",
    "parceiros": "05_parceiros.csv",
}

EMAIL_REGEX = re.compile(r"^[a-z0-9._%+-]+@(?:[a-z0-9-]+\.)+[a-z]{2,}$")


def mask_if_personal(email: str) -> str:
    """Mascara e-mails pessoais (PII/LGPD) em relatórios caso existam."""
    em = email.strip().lower()
    if "@" not in em:
        return em[:30]
    local, dom = em.split("@", 1)
    if dom in {"gmail.com", "hotmail.com", "outlook.com", "yahoo.com.br", "yahoo.com", "uol.com.br", "terra.com.br"}:
        masked_local = local[:3] + "***" if len(local) > 3 else "***"
        return f"{masked_local}@{dom}"
    return em


def load_legacy_emails_and_bounces(doc_dir: Path) -> tuple[set[str], set[str]]:
    """Extrai todos os e-mails legados de doc/*.xlsx e separa a blocklist oficial da aba 'bounce'."""
    legacy_emails: set[str] = set()
    bounce_emails: set[str] = set()

    for xlsx_path in sorted(doc_dir.glob("*.xlsx")):
        wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            is_bounce = sheet_name.strip().lower() == "bounce"
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if not cell:
                        continue
                    text = str(cell).replace("\xad", "")
                    text = re.sub(r"@(ene|mat)\.unb(?![.a-zA-Z])", r"@\1.unb.br", text, flags=re.IGNORECASE)
                    text = re.sub(r"@icex\.ufmg\.b(?![a-zA-Z])", "@icex.ufmg.br", text, flags=re.IGNORECASE)
                    for match in re.findall(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+", text):
                        cleaned = match.lower().strip(".,;:")
                        if is_bounce:
                            bounce_emails.add(cleaned)
                        else:
                            legacy_emails.add(cleaned)
        wb.close()

    return legacy_emails, bounce_emails


def check_domain_dns(domain: str, timeout: float = 3.0) -> str:
    """Verifica registros MX (ou fallback A) usando DNS públicos confiáveis (1.1.1.1 / 8.8.8.8)."""
    resolver = dns.resolver.Resolver(configure=False)
    resolver.nameservers = ["1.1.1.1", "8.8.8.8"]
    resolver.lifetime = timeout
    resolver.timeout = 1.5
    try:
        mx = resolver.resolve(domain, "MX")
        if len(mx) > 0:
            return "OK_MX"
    except dns.resolver.NXDOMAIN:
        return "NXDOMAIN"
    except Exception:
        pass
    try:
        a = resolver.resolve(domain, "A")
        if len(a) > 0:
            return "ONLY_A"
    except dns.resolver.NXDOMAIN:
        return "NXDOMAIN"
    except Exception:
        return "TIMEOUT_OR_DEAD"
    return "DEAD"


def resolve_domains_cached(domains: set[str], cache_file: Path) -> dict[str, str]:
    """Consulta DNS em paralelo com cache persistente em disco."""
    cache: dict[str, str] = {}
    if cache_file.exists():
        try:
            cache = json.loads(cache_file.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    missing = sorted(d for d in domains if d and (d not in cache or cache[d] == "TIMEOUT_OR_DEAD"))
    if missing:
        with ThreadPoolExecutor(max_workers=20) as pool:
            futures = {pool.submit(check_domain_dns, dom): dom for dom in missing}
            for fut in as_completed(futures):
                dom = futures[fut]
                try:
                    cache[dom] = fut.result()
                except Exception:
                    cache[dom] = "ERROR"
        cache_file.parent.mkdir(parents=True, exist_ok=True)
        cache_file.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")

    return {d: cache.get(d, "UNKNOWN") for d in domains}


def detect_synthetic_followers(rows: list[dict[str, str]]) -> dict[str, Any]:
    """Detecta sequências aritméticas artificiais (ex.: -100 repetido) e múltiplos redondos suspeitos."""
    nums: list[int] = []
    for r in rows:
        raw = (r.get("quantidade_seguidores") or "").strip().replace(".", "").replace(",", "")
        if raw.isdigit():
            nums.append(int(raw))
    if len(nums) < 6:
        return {"suspicious": False, "max_arithmetic_streak": 0, "round_hundred_ratio": 0.0}

    diffs = [nums[i + 1] - nums[i] for i in range(len(nums) - 1)]
    max_streak = 1
    cur_streak = 1
    for i in range(1, len(diffs)):
        if diffs[i] == diffs[i - 1] and diffs[i] != 0:
            cur_streak += 1
            max_streak = max(max_streak, cur_streak)
        else:
            cur_streak = 1

    round_hundreds = sum(1 for n in nums if n % 100 == 0)
    ratio = round(round_hundreds / len(nums), 3)
    return {
        "suspicious": max_streak >= 4 or ratio >= 0.90,
        "max_arithmetic_streak": max_streak,
        "round_hundred_ratio": ratio,
    }


def run_audit(project_root: Path, target_dir: Path | None = None) -> dict[str, Any]:
    """Executa auditoria completa nas tabelas exportadas ou em staging."""
    csv_dir = target_dir or (project_root / "csv_exports")
    doc_dir = project_root / "doc"
    verified_dir = project_root / "data_verified"
    reports_dir = project_root / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    legacy_emails, bounce_emails = load_legacy_emails_and_bounces(doc_dir)

    verified_ies_emails: set[str] = set()
    ies_file = verified_dir / "ies_scraped.json"
    if ies_file.exists():
        for item in json.loads(ies_file.read_text(encoding="utf-8")):
            if item.get("source_url") and item.get("email"):
                verified_ies_emails.add(item["email"].strip().lower())

    verified_social_handles: set[str] = set()
    social_file = verified_dir / "social_verified.json"
    if social_file.exists():
        for item in json.loads(social_file.read_text(encoding="utf-8")):
            if item.get("source_url") and item.get("arroba"):
                verified_social_handles.add(item["arroba"].strip().lower())

    all_domains: set[str] = set()
    loaded_tables: dict[str, list[dict[str, str]]] = {}
    schema_checks: dict[str, dict[str, Any]] = {}

    for key, fname in CSV_FILE_MAP.items():
        fpath = csv_dir / fname
        if not fpath.exists():
            continue
        with fpath.open(encoding="utf-8-sig", newline="") as fp:
            reader = csv.DictReader(fp)
            cols = list(reader.fieldnames or [])
            rows = list(reader)
        loaded_tables[key] = rows
        expected_key = "universidades" if key == "universidades_df" else key
        expected_cols = OFFICIAL_SCHEMAS[expected_key]
        schema_checks[key] = {
            "file": fname,
            "rows": len(rows),
            "columns": cols,
            "schema_ok": cols == expected_cols,
        }
        for r in rows:
            em = (r.get("email") or "").strip().lower()
            if "@" in em and "\n" not in em and "/" not in em and " " not in em:
                all_domains.add(em.split("@")[-1])

    dns_results = resolve_domains_cached(all_domains, verified_dir / "dns_cache.json")

    findings: dict[str, Any] = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "target_dir": str(csv_dir),
        "legacy_unique_emails_count": len(legacy_emails),
        "bounce_blocklist_count": len(bounce_emails),
        "tables": {},
        "total_blockers": 0,
    }

    total_blockers = 0

    for key, rows in loaded_tables.items():
        bad_emails: list[str] = []
        dead_domains: list[str] = []
        bounced_present: list[str] = []
        unverified_emails: list[str] = []
        dirty_depts: list[str] = []
        invalid_ufs: Counter[str] = Counter()
        float_phones: int = 0
        unverified_handles: list[str] = []

        allow_br_uf = key in {"parceiros", "midias"}
        for r in rows:
            uf = (r.get("uf") or "").strip().upper()
            if uf not in VALID_UFS and not (allow_br_uf and uf == "BR"):
                invalid_ufs[r.get("uf") or "<vazio>"] += 1

            if "nome_departamento" in r:
                dept = (r.get("nome_departamento") or "").strip()
                if (
                    not dept
                    or dept.lower() in {"nan", "none", "null", "setor específico", "departamento (outros)"}
                    or "unnamed:" in dept.lower()
                ):
                    dirty_depts.append(dept or "<vazio>")

            if "telefone" in r:
                tel = (r.get("telefone") or "").strip()
                if tel.endswith(".0"):
                    float_phones += 1

            if "email" in r:
                raw_em = r.get("email") or ""
                em = raw_em.strip().lower()
                if em:
                    has_bad_chars = ("\xad" in raw_em) or ("\n" in raw_em) or ("/" in raw_em)
                    syntax_ok = bool(EMAIL_REGEX.match(em)) and not em.endswith(".unb") and not em.endswith(".b")
                    if has_bad_chars or not syntax_ok:
                        bad_emails.append(mask_if_personal(raw_em.replace("\n", " | ")))
                    else:
                        dom = em.split("@")[-1]
                        if dns_results.get(dom) in {"NXDOMAIN", "TIMEOUT_OR_DEAD", "DEAD"}:
                            dead_domains.append(f"{mask_if_personal(em)} [{dns_results.get(dom)}]")
                    if em in bounce_emails:
                        bounced_present.append(mask_if_personal(em))
                    if key in {"universidades", "universidades_df"}:
                        if em not in legacy_emails and em not in verified_ies_emails:
                            unverified_emails.append(mask_if_personal(em))

            if "arroba" in r:
                handle = (r.get("arroba") or "").strip().lower()
                if handle and handle not in verified_social_handles:
                    unverified_handles.append(handle)

        follower_check = (
            detect_synthetic_followers(rows)
            if key in {"artistas", "influenciadores"}
            else {"suspicious": False}
        )

        table_blockers = (
            len(bad_emails)
            + len(dead_domains)
            + len(bounced_present)
            + len(unverified_emails)
            + len(dirty_depts)
            + sum(invalid_ufs.values())
            + float_phones
            + len(unverified_handles)
            + (1 if follower_check.get("suspicious") else 0)
        )
        total_blockers += table_blockers

        findings["tables"][key] = {
            **schema_checks[key],
            "bad_syntax_or_multi_emails": len(bad_emails),
            "bad_emails_sample": bad_emails[:10],
            "dead_dns_emails": len(dead_domains),
            "dead_dns_sample": dead_domains[:10],
            "bounced_emails_leaked": len(bounced_present),
            "bounced_sample": bounced_present[:10],
            "unverified_or_synthetic_emails": len(unverified_emails),
            "unverified_emails_sample": unverified_emails[:10],
            "dirty_departments_count": len(dirty_depts),
            "invalid_ufs": dict(invalid_ufs),
            "float_phones_count": float_phones,
            "unverified_social_handles": len(unverified_handles),
            "synthetic_followers_check": follower_check,
            "blockers_count": table_blockers,
        }

    findings["total_blockers"] = total_blockers
    (reports_dir / "audit_summary.json").write_text(
        json.dumps(findings, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    md_lines = [
        "# Relatório de Auditoria Técnica & Veracidade — Projeto Edumi",
        "",
        f"- **Diretório auditado:** `{csv_dir.name}`",
        f"- **E-mails únicos no acervo legado (`/doc`):** `{len(legacy_emails)}`",
        f"- **Blocklist oficial de Bounces (`/doc`):** `{len(bounce_emails)}`",
        f"- **Total de bloqueadores de qualidade encontrados:** `{total_blockers}`",
        "",
        "## Resumo por Tabela",
        "",
        "| Tabela | Linhas | Schema OK | Sintaxe/Multi-Email | Domínio Morto (DNS) | Bounces Vazados | Sem Proveniência / IA | Depto Sujo (`nan`/`Unnamed`) | UF Inválida |",
        "| :--- | ---: | :---: | ---: | ---: | ---: | ---: | ---: | :--- |",
    ]
    for key, info in findings["tables"].items():
        unver = info["unverified_or_synthetic_emails"] or info["unverified_social_handles"]
        md_lines.append(
            f"| `{info['file']}` | {info['rows']} | {'✅' if info['schema_ok'] else '❌'} "
            f"| {info['bad_syntax_or_multi_emails']} | {info['dead_dns_emails']} "
            f"| {info['bounced_emails_leaked']} | {unver} "
            f"| {info['dirty_departments_count']} | `{info['invalid_ufs'] or 'OK'}` |"
        )

    (reports_dir / "AUDITORIA_QUALIDADE.md").write_text("\n".join(md_lines) + "\n", encoding="utf-8")
    return findings
