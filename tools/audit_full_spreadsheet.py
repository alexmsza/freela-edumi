"""Script de Auditoria 360° de Todas as Abas da Planilha Centralizadora Edumi."""

import json
import re
import openpyxl
from pathlib import Path

VALID_UFS = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA",
    "MT", "MS", "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN",
    "RS", "RO", "RR", "SC", "SP", "SE", "TO"
}

EMAIL_RE = re.compile(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$")
HANDLE_RE = re.compile(r"^@[a-zA-Z0-9._]+$")

FORBIDDEN_DEPT_KEYWORDS = [
    "ouvidoria", "reitoria", "pró-reitoria", "pro-reitoria",
    "pró reitoria", "pro reitoria", "teleatendimento", "financeiro",
    "comercial", "secretaria geral", "gabinete do reitor"
]

def run_full_audit(workbook_path: str = "0. Planilha centralizadora.xlsx"):
    wb_file = Path(workbook_path)
    if not wb_file.exists():
        print(f"ERRO: Arquivo {workbook_path} não encontrado.")
        return

    wb = openpyxl.load_workbook(wb_file, data_only=True)
    report = {
        "workbook": workbook_path,
        "sheets": {},
        "global_status": "PASS",
        "total_rows": 0,
        "critical_issues": []
    }

    print(f"=== INICIANDO AUDITORIA COMPLETA: {workbook_path} ===\n")

    # 1. ABA UNIVERSIDADES
    if "universidades" in wb.sheetnames:
        ws = wb["universidades"]
        headers = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        
        uni_issues = []
        uf_counts = {}
        emails_seen = set()
        dept_violations = []

        for idx, r in enumerate(rows, start=2):
            faculdade, depto, email, uf = (r[0] or "", r[1] or "", r[2] or "", r[3] or "")
            faculdade = str(faculdade).strip()
            depto = str(depto).strip()
            email = str(email).strip().lower()
            uf = str(uf).strip().upper()

            if not faculdade:
                uni_issues.append(f"Linha {idx}: Nome de faculdade vazio")
            if not depto:
                uni_issues.append(f"Linha {idx}: Nome de departamento vazio")
            if not email or not EMAIL_RE.match(email):
                uni_issues.append(f"Linha {idx}: E-mail inválido ou vazio ({email})")
            if uf not in VALID_UFS:
                uni_issues.append(f"Linha {idx}: UF inválida ({uf})")
            
            # Checar palavras proibidas
            depto_lower = depto.lower()
            for kw in FORBIDDEN_DEPT_KEYWORDS:
                if kw in depto_lower:
                    dept_violations.append(f"Linha {idx}: Departamento com termo proibido '{kw}': {depto}")

            # Contagem de UF
            uf_counts[uf] = uf_counts.get(uf, 0) + 1
            emails_seen.add(email)

        report["sheets"]["universidades"] = {
            "total_rows": len(rows),
            "unique_emails": len(emails_seen),
            "ufs_cobertas": len(uf_counts),
            "uf_distribuicao": uf_counts,
            "issues_count": len(uni_issues),
            "dept_violations_count": len(dept_violations),
            "sample_issues": uni_issues[:5],
            "dept_violations": dept_violations[:5]
        }
        report["total_rows"] += len(rows)
        print(f"✓ universidades: {len(rows)} linhas | {len(uf_counts)} UFs cobertas | {len(dept_violations)} violações de departamento")

    # 2. ABA ARTISTAS
    if "artistas" in wb.sheetnames:
        ws = wb["artistas"]
        headers = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        art_issues = []
        handles_seen = set()
        non_df_artists = []
        banned_found = []
        ban_list = {"@ilustralu", "@beluga", "@rogercruzart", "@vitorcafaggi", "@sica_rafael"}

        for idx, r in enumerate(rows, start=2):
            nome, arroba, fol, tipo, uf = (r[0] or "", r[1] or "", r[2] or 0, r[3] or "", r[4] or "")
            nome = str(nome).strip()
            arroba = str(arroba).strip()
            uf = str(uf).strip().upper()

            if not nome:
                art_issues.append(f"Linha {idx}: Nome de artista vazio")
            if not arroba or not HANDLE_RE.match(arroba):
                art_issues.append(f"Linha {idx}: Arroba inválido ({arroba})")
            
            try:
                fol_int = int(fol)
                if fol_int <= 0:
                    art_issues.append(f"Linha {idx}: Seguidores zerados ou negativos ({fol})")
            except (ValueError, TypeError):
                art_issues.append(f"Linha {idx}: Seguidores não é número ({fol})")

            if uf != "DF":
                non_df_artists.append(f"Linha {idx}: Artista com UF diferente de DF ({uf}): {nome} ({arroba})")

            arr_lower = arroba.lower()
            if arr_lower in ban_list:
                banned_found.append(f"Linha {idx}: Perfil banido encontrado: {arroba}")

            handles_seen.add(arr_lower)

        report["sheets"]["artistas"] = {
            "total_rows": len(rows),
            "unique_handles": len(handles_seen),
            "non_df_count": len(non_df_artists),
            "banned_found_count": len(banned_found),
            "issues_count": len(art_issues),
            "sample_issues": art_issues[:5]
        }
        report["total_rows"] += len(rows)
        print(f"✓ artistas: {len(rows)} linhas | 100% DF: {len(non_df_artists) == 0} | Banidos: {len(banned_found)} | Erros: {len(art_issues)}")

    # 3. ABA INFLUENCIADORES
    if "influenciadores" in wb.sheetnames:
        ws = wb["influenciadores"]
        headers = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        inf_issues = []
        handles_seen = set()
        invalid_ufs = []
        mega_found = []
        uf_dist = {}
        mega_ban = {"@hugogloss", "@marimaria", "@hungria_oficial", "@nathaliavalente", "@sarahandrade"}

        for idx, r in enumerate(rows, start=2):
            nome, arroba, fol, tipo, uf = (r[0] or "", r[1] or "", r[2] or 0, r[3] or "", r[4] or "")
            nome = str(nome).strip()
            arroba = str(arroba).strip()
            uf = str(uf).strip().upper()

            if not nome:
                inf_issues.append(f"Linha {idx}: Nome vazio")
            if not arroba or not HANDLE_RE.match(arroba):
                inf_issues.append(f"Linha {idx}: Arroba inválido ({arroba})")
            
            try:
                fol_int = int(fol)
                if fol_int > 250000:
                    mega_found.append(f"Linha {idx}: Mega-influenciador acima de 250k: {nome} ({arroba}) com {fol_int:,} seg")
            except (ValueError, TypeError):
                inf_issues.append(f"Linha {idx}: Seguidores inválido ({fol})")

            if uf not in {"DF", "GO", "PE"}:
                invalid_ufs.append(f"Linha {idx}: UF fora do target DF/GO/PE ({uf}): {nome}")

            arr_lower = arroba.lower()
            if arr_lower in mega_ban:
                inf_issues.append(f"Linha {idx}: Mega celebridade banida presente: {arroba}")

            uf_dist[uf] = uf_dist.get(uf, 0) + 1
            handles_seen.add(arr_lower)

        report["sheets"]["influenciadores"] = {
            "total_rows": len(rows),
            "unique_handles": len(handles_seen),
            "uf_distribuicao": uf_dist,
            "invalid_ufs_count": len(invalid_ufs),
            "mega_above_250k_count": len(mega_found),
            "issues_count": len(inf_issues),
            "sample_issues": inf_issues[:5]
        }
        report["total_rows"] += len(rows)
        print(f"✓ influenciadores: {len(rows)} linhas | Distribuição: {uf_dist} | Fora de DF/GO/PE: {len(invalid_ufs)} | Erros: {len(inf_issues)}")

    # 4. ABA MIDIAS
    if "midias" in wb.sheetnames:
        ws = wb["midias"]
        headers = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        midia_issues = []
        for idx, r in enumerate(rows, start=2):
            nome = str(r[0] or "").strip()
            if not nome:
                midia_issues.append(f"Linha {idx}: Nome de mídia vazio")

        report["sheets"]["midias"] = {
            "total_rows": len(rows),
            "issues_count": len(midia_issues)
        }
        report["total_rows"] += len(rows)
        print(f"✓ midias: {len(rows)} linhas | Erros: {len(midia_issues)}")

    # 5. ABA PARCEIROS
    if "parceiros" in wb.sheetnames:
        ws = wb["parceiros"]
        headers = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        parc_issues = []
        for idx, r in enumerate(rows, start=2):
            nome = str(r[0] or "").strip()
            if not nome:
                parc_issues.append(f"Linha {idx}: Nome de parceiro vazio")

        report["sheets"]["parceiros"] = {
            "total_rows": len(rows),
            "issues_count": len(parc_issues)
        }
        report["total_rows"] += len(rows)
        print(f"✓ parceiros: {len(rows)} linhas | Erros: {len(parc_issues)}")

    print(f"\n==========================================")
    print(f"TOTAL GERAL DE REGISTROS AUDITADOS: {report['total_rows']:,}")
    print(f"==========================================")

    # Save audit report to JSON
    out_path = Path("scratch/audit_report_full.json")
    out_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Relatório detalhado salvo em {out_path}")

if __name__ == "__main__":
    import sys
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    run_full_audit("0. Planilha centralizadora.xlsx")
