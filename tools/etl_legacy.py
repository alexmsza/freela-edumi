"""Extrator determinístico e higienizador das 7 planilhas legadas em doc/*.xlsx."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any

import openpyxl

from tools.audit import EMAIL_REGEX, load_legacy_emails_and_bounces, resolve_domains_cached
from tools.uf_catalog import UNB_ACRONYMS, VALID_UFS, normalize_department, resolve_uf

EMAIL_FIND_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")

# Domínios e e-mails de Empresas Juniores que devem ir para `parceiros`
EJ_DOMAINS: set[str] = {
    "admconsultoria.com.br",
    "eletrojr.com.br",
    "engeproconsultoria.com.br",
    "mecajun.com",
    "orcestraej.com.br",
}
EJ_EMAILS: set[str] = {
    "unbflorestal@gmail.com",
}

# Mapeamento explícito dos Centros Acadêmicos e setores específicos da FT na UnB (siglas unb_2025)
UNB_SPECIFIC_EMAIL_DEPTS: dict[str, str] = {
    "saapaft@unb.br": "Serviço de Apoio Acadêmico e Pedagógico — Faculdade de Tecnologia (SAAP/FT)",
    "auditorioft@unb.br": "Coordenação e Reserva de Auditórios — Faculdade de Tecnologia (FT)",
    "dasilva@unb.br": "Secretaria Administrativa — Faculdade de Tecnologia (FT)",
    "adrianaflavia@unb.br": "Apoio Acadêmico e Administrativo — Faculdade de Tecnologia (FT)",
    "jacksonferreira@unb.br": "Secretaria de Apoio Técnico — Faculdade de Tecnologia (FT)",
    "engenhariambientalunb@gmail.com": "Centro Acadêmico de Engenharia Ambiental (CAEA/FT)",
    "unbcaenc@gmail.com": "Centro Acadêmico de Engenharia Civil (CAENC/FT)",
    "caecunb@gmail.com": "Centro Acadêmico de Engenharia de Computação (CAEC)",
    "caene@ene.unb.br": "Centro Acadêmico de Engenharia Elétrica (CAENE/FT)",
    "caef@unb.br": "Centro Acadêmico de Engenharia Florestal (CAEF/FT)",
    "camecunb@gmail.com": "Centro Acadêmico de Engenharia Mecânica (CAMEC/FT)",
    "catron.unb@gmail.com": "Centro Acadêmico de Engenharia Mecatrônica (CATRON/FT)",
    "caepro.unb@gmail.com": "Centro Acadêmico de Engenharia de Produção (CAEPRO/FT)",
    "care.unb@gmail.com": "Centro Acadêmico de Engenharia de Redes de Comunicação (CARE/FT)",
    "alunosene@ene.unb.br": "Coordenação Discente — Departamento de Engenharia Elétrica (ENE)",
    "secene@ene.unb.br": "Secretaria do Departamento de Engenharia Elétrica (ENE)",
    "secretaria@mat.unb.br": "Secretaria do Departamento de Matemática (MAT)",
    "secretaria@unb.br": "Secretaria Geral / Atendimento Institucional",
}


def clean_phone(raw: Any) -> str:
    """Remove sufixo float '.0' do Excel e normaliza espaçamento."""
    if raw is None:
        return ""
    text = str(raw).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    if re.match(r"^\d+\.0$", text):
        text = text[:-2]
    return re.sub(r"\s+", " ", text).strip()


def split_valid_emails(raw: Any, bounce_set: set[str]) -> list[str]:
    """Extrai 1 ou mais e-mails válidos de uma célula suja, corrige '.unb' truncado e filtra bounces."""
    if raw is None:
        return []
    text = str(raw).replace("\xad", "").strip()
    if not text or text.lower() in {"nan", "none", "só pelo site"}:
        return []

    text = re.sub(r"@(ene|mat)\.unb(?![.a-zA-Z])", r"@\1.unb.br", text, flags=re.IGNORECASE)
    text = re.sub(r"@icex\.ufmg\.b(?![a-zA-Z])", "@icex.ufmg.br", text, flags=re.IGNORECASE)

    found: list[str] = []
    for match in EMAIL_FIND_RE.findall(text):
        em = match.lower().strip(".,;:")
        if em in bounce_set:
            continue
        if EMAIL_REGEX.match(em) and not re.match(r"^[0-9a-f]{20,}@", em):
            if em not in found:
                found.append(em)
    return found


def clean_institution_name(raw_name: str, email: str) -> tuple[str, str]:
    """Separa 'Sigla - Setor' quando colados na coluna nome e padroniza nome da instituição."""
    name = re.sub(r"\s+", " ", (raw_name or "").strip())
    inferred_dept = ""

    if email.endswith("@unb.br") or ".unb.br" in email:
        if name.upper() in UNB_ACRONYMS:
            return "Universidade de Brasília (UnB)", UNB_ACRONYMS[name.upper()]
        if name.upper().startswith("UNB - "):
            sub = name.split("-", 1)[1].strip()
            return "Universidade de Brasília (UnB)", UNB_ACRONYMS.get(sub.upper(), sub)
        if name.upper() == "UNB" or not name:
            return "Universidade de Brasília (UnB)", ""

    if email.endswith("@iesb.br") or name.upper() == "IESB":
        return "Centro Universitário IESB", inferred_dept
    if email.endswith("@ucb.br") and name.upper() in {"UCB", "UNIVERSIDADE CATÓLICA DE BRASÍLIA"}:
        return "Universidade Católica de Brasília (UCB)", inferred_dept
    if email.endswith("@uniceub.br") and name.upper() in {"CEUB", "UNICEUB"}:
        return "Centro Universitário de Brasília (CEUB / UniCEUB)", inferred_dept
    if email.endswith("@udf.edu.br") and name.upper() == "UDF":
        return "Centro Universitário do Distrito Federal (UDF)", inferred_dept

    if " - " in name and len(name) < 60:
        left, right = [p.strip() for p in name.split(" - ", 1)]
        if left.isupper() and len(left) <= 10:
            return left, right

    return name or "Instituição Acadêmica", inferred_dept


def ingest_legacy_spreadsheets(project_root: Path, out_dir: Path | None = None) -> dict[str, int]:
    """Extrai, desduplica, filtra domínios mortos (DNS) e higieniza todas as bases legadas de doc/."""
    doc_dir = project_root / "doc"
    staged_dir = out_dir or (project_root / "data_staged")
    verified_dir = project_root / "data_verified"
    reports_dir = project_root / "reports"
    staged_dir.mkdir(parents=True, exist_ok=True)
    verified_dir.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)

    _, bounce_emails = load_legacy_emails_and_bounces(doc_dir)

    uni_by_email: dict[str, dict[str, str]] = {}
    midias_rows: list[dict[str, str]] = []
    parceiros_by_key: dict[str, dict[str, str]] = {}
    provenance_log: list[dict[str, str]] = []

    def upsert_uni(fac: str, dept: str, email: str, raw_uf: str | None, source: str) -> None:
        dom = email.split("@")[-1]
        if dom in EJ_DOMAINS or email in EJ_EMAILS:
            if dom != "orcestraej.com.br":
                parceiros_by_key[email] = {
                    "nome": f"Empresa Júnior ({fac} / {dept or email.split('@')[0]})",
                    "link": f"https://{dom}" if dom != "gmail.com" else "",
                    "email": email,
                    "telefone": "",
                    "uf": "DF",
                    "obs": f"Empresa Júnior universitária extraída de {source}",
                }
            return

        fac_clean, sub_dept = clean_institution_name(fac, email)
        if email in UNB_SPECIFIC_EMAIL_DEPTS:
            fac_clean = "Universidade de Brasília (UnB)"
            dept_clean = UNB_SPECIFIC_EMAIL_DEPTS[email]
        else:
            dept_candidate = dept or sub_dept
            dept_clean = normalize_department(dept_candidate, email=email, inst_name=fac_clean)

        uf_clean = resolve_uf(raw_uf, email=email, institution_name=f"{fac_clean} {dept_clean}")
        if uf_clean not in VALID_UFS:
            return

        existing = uni_by_email.get(email)
        if existing:
            if "Atendimento Institucional" in existing["nome_departamento"] and "Atendimento Institucional" not in dept_clean:
                existing["nome_departamento"] = dept_clean
            if len(fac_clean) > len(existing["nome_faculdade"]):
                existing["nome_faculdade"] = fac_clean
            return

        uni_by_email[email] = {
            "nome_faculdade": fac_clean,
            "nome_departamento": dept_clean,
            "email": email,
            "uf": uf_clean,
        }
        provenance_log.append({"tab": "universidades", "key": email, "source": source})

    # 1. Audiencias faculdades.xlsx
    f_aud = doc_dir / "Audiencias faculdades.xlsx"
    if f_aud.exists():
        wb = openpyxl.load_workbook(f_aud, data_only=True)
        unb_sigla_map: dict[str, str] = {}
        if "unb_2025" in wb.sheetnames:
            for idx, r in enumerate(wb["unb_2025"].iter_rows(values_only=True), start=1):
                if idx == 1 or not r or not any(r):
                    continue
                local, sigla, raw_em = (r[0], r[1], r[2]) if len(r) >= 3 else (r[0], "", None)
                local_s = str(local or "").strip()
                sigla_s = str(sigla or "").strip()
                if sigla_s and local_s and local_s.upper() != sigla_s.upper():
                    unb_sigla_map[sigla_s.upper()] = local_s
                for em in split_valid_emails(raw_em, bounce_emails):
                    if sigla_s.upper() in UNB_ACRONYMS and (not local_s or local_s.upper() == sigla_s.upper()):
                        dept_label = UNB_ACRONYMS[sigla_s.upper()]
                    elif local_s and sigla_s and local_s.upper() != sigla_s.upper():
                        dept_label = f"{local_s} ({sigla_s})"
                    else:
                        dept_label = local_s or sigla_s
                    upsert_uni("Universidade de Brasília (UnB)", dept_label, em, "DF", "Audiencias:unb_2025")

        if "siglas unb_2025" in wb.sheetnames:
            last_sigla = ""
            for idx, r in enumerate(wb["siglas unb_2025"].iter_rows(values_only=True), start=1):
                if idx == 1 or not r or not any(r):
                    continue
                sigla = str(r[0] or "").strip()
                if sigla:
                    last_sigla = sigla
                else:
                    sigla = last_sigla
                raw_em = r[1] if len(r) > 1 else None
                dept_label = unb_sigla_map.get(sigla.upper(), UNB_ACRONYMS.get(sigla.upper(), sigla))
                for em in split_valid_emails(raw_em, bounce_emails):
                    upsert_uni("Universidade de Brasília (UnB)", dept_label, em, "DF", "Audiencias:siglas_unb_2025")

        for sname in ("faculdades_e_ifs_1_onda", "heranca_paises_baixos_2025"):
            if sname in wb.sheetnames:
                for idx, r in enumerate(wb[sname].iter_rows(values_only=True), start=1):
                    if idx == 1 or not r or not any(r):
                        continue
                    nome = str(r[0] or "").strip()
                    raw_em = r[1] if len(r) > 1 else None
                    uf_col = str(r[2] or "").strip() if len(r) > 2 and r[2] else None
                    for em in split_valid_emails(raw_em, bounce_emails):
                        if sname == "heranca_paises_baixos_2025":
                            upsert_uni("", nome, em, uf_col, f"Audiencias:{sname}")
                        else:
                            upsert_uni(nome, "", em, uf_col, f"Audiencias:{sname}")
        wb.close()

    # 2. Goiás V4
    for f_go in doc_dir.glob("*Goia*V4.xlsx"):
        wb = openpyxl.load_workbook(f_go, data_only=True)
        for sname in wb.sheetnames:
            for idx, r in enumerate(wb[sname].iter_rows(values_only=True), start=1):
                if idx == 1 or not r or not any(r):
                    continue
                inst = str(r[0] or "").strip()
                cidade = str(r[2] or "").strip() if len(r) > 2 and r[2] else ""
                raw_em = r[4] if len(r) > 4 else None
                obs = str(r[5] or "").strip() if len(r) > 5 and r[5] else ""
                dept_hint = obs if (obs and len(obs) <= 70) else (f"Unidade {cidade}" if cidade and cidade != "Várias Cidades" else "")
                for em in split_valid_emails(raw_em, bounce_emails):
                    upsert_uni(inst, dept_hint, em, "GO", f"GoiasV4:{sname}")
        wb.close()

    # 3. Lista de Psicologia
    for f_psi in doc_dir.glob("Lista de Psicologia*.xlsx"):
        wb = openpyxl.load_workbook(f_psi, data_only=True)
        for sname in wb.sheetnames:
            for idx, r in enumerate(wb[sname].iter_rows(values_only=True), start=1):
                if idx == 1 or not r or not any(r):
                    continue
                raw_nome = str(r[0] or "").strip()
                raw_em = r[1] if len(r) > 1 else None
                uf_col = str(r[2] or "").strip() if len(r) > 2 and r[2] else None
                m = re.search(r"\bda\s+(.+)$", raw_nome, flags=re.IGNORECASE)
                inst_name = m.group(1).strip() if m else raw_nome
                for em in split_valid_emails(raw_em, bounce_emails):
                    upsert_uni(inst_name, "Departamento / Coordenação de Psicologia", em, uf_col, f"Psicologia:{sname}")
        wb.close()

    # 4. Lista de e-mails Universidades Brasil.xlsx & arquivado.xlsx
    for fname in ("Lista de e-mails Universidades Brasil.xlsx", "arquivado.xlsx"):
        fpath = doc_dir / fname
        if not fpath.exists():
            continue
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for sname in wb.sheetnames:
            if sname.strip().lower() == "bounce":
                continue
            rows = list(wb[sname].iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(h or "").strip() for h in rows[0]]
            for r in rows[1:]:
                if not r or not any(r):
                    continue
                inst = str(r[0] or "").strip()
                if len(headers) >= 2 and headers[1].lower() in {"email", "e-mail", "e-mail´s"}:
                    uf_val = str(r[2] or "").strip() if len(headers) >= 3 and headers[2].lower() == "estado" and len(r) > 2 else None
                    for col_idx in range(1, len(r)):
                        if col_idx == 2 and uf_val:
                            continue
                        for em in split_valid_emails(r[col_idx], bounce_emails):
                            upsert_uni(inst, "", em, uf_val, f"{fname}:{sname}")
                else:
                    uf_from_sheet = None
                    if "RJ" in sname:
                        uf_from_sheet = "RJ"
                    elif "SP" in sname:
                        uf_from_sheet = "SP"
                    elif "MG" in sname:
                        uf_from_sheet = "MG"
                    for col_idx in range(1, len(r)):
                        col_h = headers[col_idx] if col_idx < len(headers) else ""
                        if col_h.upper() == "SITE":
                            continue
                        dept_h = col_h if col_h and col_h.lower() not in {"geral", "outros"} else ""
                        for em in split_valid_emails(r[col_idx], bounce_emails):
                            upsert_uni(inst, dept_h, em, uf_from_sheet, f"{fname}:{sname}")
        wb.close()

    # 5. Mídias (jornal).xlsx
    for f_mid in doc_dir.glob("M*dias*.xlsx"):
        wb = openpyxl.load_workbook(f_mid, data_only=True)
        for sname in wb.sheetnames:
            uf_sheet = resolve_uf(sname) or "BR"
            rows = [r for r in wb[sname].iter_rows(values_only=True) if r and any(x is not None and str(x).strip() != "" for x in r)]
            if not rows:
                continue
            first_col0 = str(rows[0][0] or "").strip().lower()
            has_header = first_col0 in {"sites", "nome", "site", "veículo"}
            data_rows = rows[1:] if has_header else rows

            for r in data_rows:
                nome = str(r[0] or "").strip()
                link = str(r[1] or "").strip() if len(r) > 1 and r[1] else ""
                raw_em = r[2] if len(r) > 2 else None
                tel = clean_phone(r[3]) if len(r) > 3 else ""
                if not link.startswith("http") and re.search(r"\(\d{2}\)|\d{8,}", link):
                    tel = clean_phone(link)
                    link = ""
                emails = split_valid_emails(raw_em, bounce_emails)
                if "instagram.com" in link.lower():
                    key = link.lower().split("?")[0].rstrip("/")
                    parceiros_by_key[key] = {
                        "nome": nome,
                        "link": link.split("?")[0],
                        "email": emails[0] if emails else "",
                        "telefone": tel,
                        "uf": uf_sheet,
                        "obs": f"Canal regional de divulgação ({sname})",
                    }
                elif emails:
                    for em in emails:
                        midias_rows.append({
                            "nome": nome,
                            "link": link,
                            "email": em,
                            "telefone": tel,
                            "uf": uf_sheet,
                            "obs": f"Imprensa / Portal ({sname})",
                        })
                elif nome and (link or tel):
                    midias_rows.append({
                        "nome": nome,
                        "link": link,
                        "email": "",
                        "telefone": tel,
                        "uf": uf_sheet,
                        "obs": f"Imprensa / Contato via site ou telefone ({sname})",
                    })
        wb.close()

    # 6. Contas instagram parcerias para divulgação.xlsx
    for f_parc in doc_dir.glob("Contas instagram*.xlsx"):
        wb = openpyxl.load_workbook(f_parc, data_only=True)
        for sname in wb.sheetnames:
            rows = [r for r in wb[sname].iter_rows(values_only=True) if r and any(x is not None and str(x).strip() != "" for x in r)]
            if len(rows) <= 1:
                continue
            for r in rows[1:]:
                nome = str(r[0] or "").strip()
                if not nome:
                    continue
                if sname == "zapzaps":
                    tel = clean_phone(r[1]) if len(r) > 1 else ""
                    pessoa = str(r[2] or "").strip() if len(r) > 2 and r[2] else ""
                    uf_inferred = resolve_uf(None, institution_name=nome) or "BR"
                    parceiros_by_key[f"zap:{nome.lower()}"] = {
                        "nome": nome,
                        "link": "",
                        "email": "",
                        "telefone": tel,
                        "uf": uf_inferred,
                        "obs": f"Contato direto WhatsApp ({pessoa})" if pessoa else "Contato direto WhatsApp",
                    }
                elif sname == "Instagrans com pessoas publicas":
                    link = str(r[1] or "").strip() if len(r) > 1 and r[1] else ""
                    key = (link or nome).lower().rstrip("/")
                    if key not in parceiros_by_key:
                        parceiros_by_key[key] = {
                            "nome": nome,
                            "link": link,
                            "email": "",
                            "telefone": "",
                            "uf": resolve_uf(None, institution_name=nome) or "BR",
                            "obs": "Perfil público / Parceiro de divulgação",
                        }
                else:
                    desc = str(r[1] or "").strip() if len(r) > 1 and r[1] else ""
                    link = str(r[2] or "").strip() if len(r) > 2 and r[2] else ""
                    plataforma = "TikTok" if sname == "tiktoks" else "Instagram"
                    obs_parts = [p for p in (desc, f"Canal {plataforma}") if p]
                    key = (link or nome).lower().rstrip("/")
                    parceiros_by_key[key] = {
                        "nome": nome,
                        "link": link,
                        "email": "",
                        "telefone": "",
                        "uf": resolve_uf(None, institution_name=f"{nome} {desc}") or "BR",
                        "obs": " — ".join(obs_parts)[:180],
                    }
        wb.close()

    # 7. Quarentena automática de domínios legados expirados (NXDOMAIN / DEAD)
    candidate_domains = {em.split("@")[-1] for em in uni_by_email}
    candidate_domains.update(r["email"].split("@")[-1] for r in midias_rows if "@" in r["email"])
    dns_map = resolve_domains_cached(candidate_domains, verified_dir / "dns_cache.json")

    quarantined: list[dict[str, str]] = []
    live_uni: list[dict[str, str]] = []
    for em, row in uni_by_email.items():
        dom = em.split("@")[-1]
        if dns_map.get(dom) in {"OK_MX", "ONLY_A"}:
            live_uni.append(row)
        else:
            quarantined.append({"tab": "universidades", **row, "dns_status": dns_map.get(dom, "DEAD")})

    live_midias: list[dict[str, str]] = []
    for row in midias_rows:
        em = row["email"]
        if em and dns_map.get(em.split("@")[-1]) not in {"OK_MX", "ONLY_A"}:
            quarantined.append({"tab": "midias", **row, "dns_status": dns_map.get(em.split("@")[-1], "DEAD")})
        else:
            live_midias.append(row)

    (reports_dir / "quarantined_dead_domains.json").write_text(
        json.dumps(quarantined, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    uni_list = sorted(live_uni, key=lambda x: (x["uf"] != "DF", x["uf"], x["nome_faculdade"], x["nome_departamento"]))
    uni_df_list = [u for u in uni_list if u["uf"] == "DF"]
    parceiros_list = list(parceiros_by_key.values())

    def write_csv(fname: str, fieldnames: list[str], data: list[dict[str, str]]) -> None:
        with (staged_dir / fname).open("w", encoding="utf-8-sig", newline="") as fp:
            writer = csv.DictWriter(fp, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(data)

    write_csv("01_universidades_DF_marco1.csv", ["nome_faculdade", "nome_departamento", "email", "uf"], uni_df_list)
    write_csv("01_universidades.csv", ["nome_faculdade", "nome_departamento", "email", "uf"], uni_list)
    write_csv("04_midias.csv", ["nome", "link", "email", "telefone", "uf", "obs"], live_midias)
    write_csv("05_parceiros.csv", ["nome", "link", "email", "telefone", "uf", "obs"], parceiros_list)
    for fn in ("02_artistas.csv", "03_influenciadores.csv"):
        if not (staged_dir / fn).exists():
            write_csv(fn, ["nome", "arroba", "quantidade_seguidores", "tipo", "uf"], [])

    (staged_dir / "provenance_legacy.json").write_text(
        json.dumps(provenance_log, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    return {
        "universidades_df": len(uni_df_list),
        "universidades_total": len(uni_list),
        "midias": len(live_midias),
        "parceiros": len(parceiros_list),
        "quarantined_dead_domains": len(quarantined),
    }
